"""
Simulador de chatbot para Mimbrería Fleming
Trabajo Práctico Integrador - Organización Empresarial

Este programa simula por consola la atención inicial de un chatbot.
No está integrado a WhatsApp, pero representa el mismo flujo del BPMN:
- Menú principal
- Consulta de productos en Excel
- Validación de stock
- Camino feliz
- Camino infeliz
- Registro de pedido
- Máquina de estados mediante la variable "estado"
"""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ARCHIVO_EXCEL = "Base_Datos_Mimbreria.xlsx"


def cargar_excel():
    """Carga el archivo Excel que funciona como base de datos simulada."""
    try:
        libro = load_workbook(ARCHIVO_EXCEL)
        return libro
    except FileNotFoundError:
        print("\nERROR: No se encontró el archivo Base_Datos_Mimbreria.xlsx.")
        print("Verificá que el Excel esté en la misma carpeta que este archivo Python.\n")
        return None


def obtener_productos(libro):
    """Lee la hoja Productos y devuelve una lista de productos disponibles."""
    hoja = libro["Productos"]
    productos = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[0] is not None:
            producto = {
                "nombre": str(fila[0]),
                "precio": fila[1],
                "stock": fila[2],
                "medida": str(fila[3])
            }
            productos.append(producto)

    return productos


def mostrar_menu_principal():
    """Muestra el menú inicial del chatbot."""
    print("\n===== CHATBOT MIMBRERÍA FLEMING =====")
    print("1. Consultar productos")
    print("2. Realizar pedido")
    print("3. Salir")


def mostrar_productos(productos):
    """Muestra los productos cargados en la base de datos simulada."""
    print("\nProductos disponibles en la base de datos:")
    for i, producto in enumerate(productos, start=1):
        print(f"{i}. {producto['nombre']} - Medida: {producto['medida']}")


def pedir_opcion_numerica(mensaje, minimo, maximo):
    """Valida que el usuario ingrese un número dentro del rango indicado."""
    while True:
        dato = input(mensaje).strip()

        if not dato.isdigit():
            print("Entrada inválida. Debe ingresar un número.")
            continue

        numero = int(dato)

        if numero < minimo or numero > maximo:
            print(f"Opción inválida. Ingrese un número entre {minimo} y {maximo}.")
            continue

        return numero


def pedir_si_no(mensaje):
    """Valida respuestas de tipo sí/no."""
    while True:
        respuesta = input(mensaje).strip().lower()

        if respuesta in ["si", "sí", "s"]:
            return "Sí"
        elif respuesta in ["no", "n"]:
            return "No"
        else:
            print("Respuesta inválida. Ingrese sí o no.")


def pedir_texto_obligatorio(mensaje):
    """Evita que el usuario deje datos vacíos."""
    while True:
        texto = input(mensaje).strip()

        if texto == "":
            print("Este dato no puede quedar vacío.")
        else:
            return texto


def seleccionar_producto(productos):
    """Permite elegir un producto de la lista."""
    mostrar_productos(productos)
    opcion = pedir_opcion_numerica("\nSeleccione el número del producto: ", 1, len(productos))
    return productos[opcion - 1]


def consultar_producto(producto):
    """Muestra la información del producto consultado."""
    print("\nInformación del producto:")
    print(f"Producto: {producto['nombre']}")
    print(f"Precio: ${producto['precio']}")
    print(f"Stock disponible: {producto['stock']}")
    print(f"Medida: {producto['medida']}")


def buscar_fila_producto(libro, nombre_producto):
    """Busca en qué fila del Excel está el producto elegido."""
    hoja = libro["Productos"]

    for fila in range(2, hoja.max_row + 1):
        valor = hoja.cell(row=fila, column=1).value
        if valor == nombre_producto:
            return fila

    return None


def obtener_proximo_id_pedido(libro):
    """Calcula el próximo ID de pedido según la hoja Pedidos."""
    hoja = libro["Pedidos"]
    ids = []

    for fila in range(2, hoja.max_row + 1):
        valor = hoja.cell(row=fila, column=1).value
        if isinstance(valor, int):
            ids.append(valor)

    if len(ids) == 0:
        return 1

    return max(ids) + 1


def ajustar_columnas(hoja):
    """Ajusta el ancho de las columnas para que el Excel quede más prolijo."""
    for columna in hoja.columns:
        maximo = 0
        letra = get_column_letter(columna[0].column)

        for celda in columna:
            if celda.value is not None:
                maximo = max(maximo, len(str(celda.value)))

        hoja.column_dimensions[letra].width = maximo + 2


def registrar_pedido(libro, producto, cantidad, nombre_cliente, forma_pago, envio, direccion):
    """Registra el pedido en la hoja Pedidos y descuenta el stock."""
    hoja_pedidos = libro["Pedidos"]
    hoja_productos = libro["Productos"]

    id_pedido = obtener_proximo_id_pedido(libro)

    hoja_pedidos.append([
        id_pedido,
        nombre_cliente,
        producto["nombre"],
        cantidad,
        forma_pago,
        envio,
        "Pendiente"
    ])

    fila_producto = buscar_fila_producto(libro, producto["nombre"])

    if fila_producto is not None:
        stock_actual = hoja_productos.cell(row=fila_producto, column=3).value
        hoja_productos.cell(row=fila_producto, column=3).value = stock_actual - cantidad

    ajustar_columnas(hoja_pedidos)
    ajustar_columnas(hoja_productos)

    libro.save(ARCHIVO_EXCEL)

    print("\nPedido registrado correctamente.")
    print(f"Número de pedido: {id_pedido}")
    print(f"Cliente: {nombre_cliente}")
    print(f"Producto: {producto['nombre']}")
    print(f"Cantidad: {cantidad}")
    print(f"Forma de pago: {forma_pago}")
    print(f"Envío: {envio}")

    if envio == "Sí":
        print(f"Dirección: {direccion}")

    print("Estado del pedido: Pendiente")


def flujo_pedido(libro):
    """Ejecuta el camino feliz y contempla caminos infelices."""
    estado = "consulta_producto"
    productos = obtener_productos(libro)

    if len(productos) == 0:
        print("\nNo hay productos cargados en la base de datos.")
        return

    print(f"\nEstado actual: {estado}")
    producto = seleccionar_producto(productos)

    estado = "mostrar_informacion"
    print(f"Estado actual: {estado}")
    consultar_producto(producto)

    estado = "validar_stock"
    print(f"Estado actual: {estado}")

    if producto["stock"] <= 0:
        print("\nEl producto existe, pero actualmente no tiene stock disponible.")
        print("Camino infeliz: el sistema informa el problema y vuelve al menú principal.")
        return

    estado = "confirmar_compra"
    print(f"Estado actual: {estado}")
    quiere_comprar = pedir_si_no("\n¿Desea realizar la compra? (sí/no): ")

    if quiere_comprar == "No":
        print("\nOperación cancelada por el cliente. Fin del proceso.")
        return

    estado = "solicitar_datos"
    print(f"Estado actual: {estado}")
    nombre_cliente = pedir_texto_obligatorio("Ingrese su nombre y apellido: ")

    cantidad = pedir_opcion_numerica(
        f"Ingrese cantidad a comprar. Stock disponible: {producto['stock']}: ",
        1,
        producto["stock"]
    )

    estado = "forma_pago"
    print(f"Estado actual: {estado}")
    print("\nFormas de pago:")
    print("1. Seña")
    print("2. Pago completo")
    opcion_pago = pedir_opcion_numerica("Seleccione forma de pago: ", 1, 2)

    if opcion_pago == 1:
        forma_pago = "Seña"
    else:
        forma_pago = "Pago completo"

    estado = "envio"
    print(f"Estado actual: {estado}")
    envio = pedir_si_no("¿Solicita envío? (sí/no): ")

    if envio == "Sí":
        direccion = pedir_texto_obligatorio("Ingrese dirección de entrega: ")
    else:
        direccion = "Retiro / sin envío"

    estado = "registrar_pedido"
    print(f"Estado actual: {estado}")
    registrar_pedido(libro, producto, cantidad, nombre_cliente, forma_pago, envio, direccion)

    estado = "confirmacion_final"
    print(f"Estado actual: {estado}")
    print("\nGracias por comunicarse con Mimbrería Fleming.")
    print("Un integrante del equipo continuará la gestión del pedido si es necesario.")


def main():
    """Función principal del simulador."""
    libro = cargar_excel()

    if libro is None:
        return

    estado = "menu_principal"

    while True:
        print(f"\nEstado actual: {estado}")
        mostrar_menu_principal()

        opcion = input("\nIngrese una opción: ").strip()

        if opcion == "1":
            estado = "consulta_productos"
            productos = obtener_productos(libro)

            if len(productos) == 0:
                print("\nNo hay productos cargados.")
            else:
                mostrar_productos(productos)

        elif opcion == "2":
            estado = "realizar_pedido"
            flujo_pedido(libro)
            estado = "menu_principal"

        elif opcion == "3":
            estado = "fin"
            print("\nGracias por utilizar el chatbot de Mimbrería Fleming.")
            print("Fin del proceso.")
            break

        else:
            print("\nOpción inválida. Ingrese 1, 2 o 3.")
            print("Camino infeliz: el bot detecta un input incorrecto y vuelve al menú.")


if __name__ == "__main__":
    main()
